#!/usr/bin/env python3
"""Generate minimalist blue product quarterly summary PPT template."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Minimalist blue palette
BLUE_PRIMARY = RGBColor(0x1A, 0x56, 0xDB)  # #1A56DB
BLUE_DARK = RGBColor(0x1E, 0x3A, 0x5F)  # #1E3A5F
BLUE_LIGHT = RGBColor(0xE8, 0xF0, 0xFE)  # #E8F0FE
BLUE_ACCENT = RGBColor(0x3B, 0x82, 0xF6)  # #3B82F6
GRAY_TEXT = RGBColor(0x64, 0x74, 0x8B)  # #64748B
GRAY_DARK = RGBColor(0x1E, 0x29, 0x3B)  # #1E293B
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_slide_white(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def add_top_bar(slide, prs, height=Inches(0.12)):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, height
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE_PRIMARY
    bar.line.fill.background()


def add_left_accent(slide, height=Inches(1.2)):
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.35), Inches(0.08), height
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE_ACCENT
    accent.line.fill.background()


def add_footer(slide, prs, text="产品季度总结 | 2026 Q__"):
    box = slide.shapes.add_textbox(
        Inches(0.6), prs.slide_height - Inches(0.55), Inches(6), Inches(0.35)
    )
    tf = box.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = GRAY_TEXT
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.6),
        prs.slide_height - Inches(0.65),
        prs.slide_width - Inches(1.2),
        Pt(1),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE_LIGHT
    line.line.fill.background()


def style_title(text_frame, text, size=32, color=GRAY_DARK, bold=True):
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold


def add_bullets(text_frame, items, size=18, color=GRAY_DARK, level0=True):
    text_frame.clear()
    text_frame.word_wrap = True
    for i, item in enumerate(items):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text = item
        p.level = 0 if level0 else 0
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(12)
        if not level0:
            p.bullet = True


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    # Left blue panel
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.8), prs.slide_height
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = BLUE_PRIMARY
    panel.line.fill.background()
    # Decorative circle
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(-0.8), Inches(4.2), Inches(3.2), Inches(3.2)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = BLUE_ACCENT
    circle.fill.transparency = 0.35
    circle.line.fill.background()
    # Title on white area
    title = slide.shapes.add_textbox(Inches(5.4), Inches(2.4), Inches(7.2), Inches(1.2))
    style_title(title.text_frame, "产品季度总结", size=44, color=GRAY_DARK)
    sub = slide.shapes.add_textbox(Inches(5.4), Inches(3.5), Inches(7.2), Inches(0.8))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "20__ 年第 __ 季度"
    sp.font.size = Pt(22)
    sp.font.color.rgb = GRAY_TEXT
    dept = slide.shapes.add_textbox(Inches(5.4), Inches(4.4), Inches(7.2), Inches(0.6))
    dp = dept.text_frame.paragraphs[0]
    dp.text = "汇报部门 / 汇报人"
    dp.font.size = Pt(16)
    dp.font.color.rgb = BLUE_PRIMARY
    # Vertical label on panel
    lbl = slide.shapes.add_textbox(Inches(0.9), Inches(2.8), Inches(3.5), Inches(2))
    lp = lbl.text_frame.paragraphs[0]
    lp.text = "QUARTERLY\nREVIEW"
    lp.font.size = Pt(20)
    lp.font.color.rgb = WHITE
    lp.font.bold = True
    lp.line_spacing = 1.2


def slide_section(prs, title, subtitle, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    add_top_bar(slide, prs)
    add_left_accent(slide, Inches(1.4))
    tb = slide.shapes.add_textbox(Inches(0.95), Inches(1.25), Inches(11), Inches(0.9))
    style_title(tb.text_frame, title, size=36)
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.95), Inches(2.15), Inches(10), Inches(0.5))
        sp = sb.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)
        sp.font.color.rgb = GRAY_TEXT
    add_footer(slide, prs, footer)


def slide_toc(prs, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    add_top_bar(slide, prs)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.55), Inches(4), Inches(0.7))
    style_title(tb.text_frame, "目录", size=32)
    items = [
        "01  季度概览",
        "02  核心指标与数据",
        "03  产品进展与迭代",
        "04  亮点与成果",
        "05  问题与挑战",
        "06  下季度规划",
    ]
    y = 1.35
    for i, item in enumerate(items):
        num_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6),
            Inches(y),
            Inches(0.55),
            Inches(0.55),
        )
        num_shape.fill.solid()
        num_shape.fill.fore_color.rgb = BLUE_LIGHT if i % 2 else BLUE_PRIMARY
        num_shape.line.fill.background()
        nf = num_shape.text_frame
        nf.text = f"{i + 1:02d}"
        nf.paragraphs[0].alignment = PP_ALIGN.CENTER
        nf.vertical_anchor = MSO_ANCHOR.MIDDLE
        nf.paragraphs[0].font.size = Pt(14)
        nf.paragraphs[0].font.bold = True
        nf.paragraphs[0].font.color.rgb = (
            BLUE_PRIMARY if i % 2 else WHITE
        )
        tx = slide.shapes.add_textbox(Inches(1.35), Inches(y + 0.08), Inches(8), Inches(0.45))
        tp = tx.text_frame.paragraphs[0]
        tp.text = item.split("  ", 1)[1]
        tp.font.size = Pt(20)
        tp.font.color.rgb = GRAY_DARK
        y += 0.85
    add_footer(slide, prs, footer)


def slide_overview(prs, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    add_top_bar(slide, prs)
    add_left_accent(slide)
    tb = slide.shapes.add_textbox(Inches(0.95), Inches(0.55), Inches(8), Inches(0.7))
    style_title(tb.text_frame, "季度概览", size=32)
    cards = [
        ("本季目标", "简要描述本季度产品目标与战略重点"),
        ("整体结论", "一句话总结本季度达成情况"),
        ("关键变化", "市场、用户或业务侧的重要变化"),
    ]
    x_positions = [0.6, 4.55, 8.5]
    for x, (title, desc) in enumerate(zip([c[0] for c in cards], [c[1] for c in cards])):
        left = x_positions[x]
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(1.5),
            Inches(3.7),
            Inches(4.8),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = BLUE_LIGHT
        card.line.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)
        ct = slide.shapes.add_textbox(Inches(left + 0.25), Inches(1.75), Inches(3.2), Inches(0.5))
        style_title(ct.text_frame, cards[x][0], size=20, color=BLUE_DARK)
        cd = slide.shapes.add_textbox(Inches(left + 0.25), Inches(2.35), Inches(3.2), Inches(3.5))
        cp = cd.text_frame.paragraphs[0]
        cp.text = cards[x][1]
        cp.font.size = Pt(14)
        cp.font.color.rgb = GRAY_TEXT
        cd.text_frame.word_wrap = True
    add_footer(slide, prs, footer)


def slide_kpi(prs, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    add_top_bar(slide, prs)
    add_left_accent(slide)
    tb = slide.shapes.add_textbox(Inches(0.95), Inches(0.55), Inches(8), Inches(0.7))
    style_title(tb.text_frame, "核心指标与数据", size=32)
    metrics = [
        ("DAU / MAU", "—", "环比 +__%"),
        ("留存率", "—", "目标 __%"),
        ("转化率", "—", "环比 +__%"),
        ("NPS / 满意度", "—", "较上季 +__"),
    ]
    for row, (name, value, delta) in enumerate(metrics):
        y = 1.45 + row * 1.25
        bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6),
            Inches(y),
            Inches(12.1),
            Inches(1.05),
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = WHITE if row % 2 else BLUE_LIGHT
        bg.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        nbox = slide.shapes.add_textbox(Inches(0.85), Inches(y + 0.28), Inches(3.5), Inches(0.5))
        nbox.text_frame.paragraphs[0].text = name
        nbox.text_frame.paragraphs[0].font.size = Pt(16)
        nbox.text_frame.paragraphs[0].font.color.rgb = GRAY_DARK
        vbox = slide.shapes.add_textbox(Inches(4.5), Inches(y + 0.15), Inches(2.5), Inches(0.75))
        vp = vbox.text_frame.paragraphs[0]
        vp.text = value
        vp.font.size = Pt(28)
        vp.font.bold = True
        vp.font.color.rgb = BLUE_PRIMARY
        dbox = slide.shapes.add_textbox(Inches(10), Inches(y + 0.28), Inches(2.3), Inches(0.5))
        dp = dbox.text_frame.paragraphs[0]
        dp.text = delta
        dp.font.size = Pt(14)
        dp.font.color.rgb = GRAY_TEXT
        dp.alignment = PP_ALIGN.RIGHT
    hint = slide.shapes.add_textbox(Inches(0.6), Inches(6.55), Inches(11), Inches(0.4))
    hint.text_frame.paragraphs[0].text = "提示：替换占位符为实际数据，可补充趋势小图或截图"
    hint.text_frame.paragraphs[0].font.size = Pt(11)
    hint.text_frame.paragraphs[0].font.color.rgb = GRAY_TEXT
    add_footer(slide, prs, footer)


def slide_content_list(prs, title, bullets, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    add_top_bar(slide, prs)
    add_left_accent(slide, Inches(2.5))
    tb = slide.shapes.add_textbox(Inches(0.95), Inches(0.55), Inches(10), Inches(0.7))
    style_title(tb.text_frame, title, size=32)
    body = slide.shapes.add_textbox(Inches(0.95), Inches(1.45), Inches(11.5), Inches(5.2))
    add_bullets(body.text_frame, bullets, size=20)
    add_footer(slide, prs, footer)


def slide_timeline(prs, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    add_top_bar(slide, prs)
    add_left_accent(slide)
    tb = slide.shapes.add_textbox(Inches(0.95), Inches(0.55), Inches(8), Inches(0.7))
    style_title(tb.text_frame, "产品进展与迭代", size=32)
    months = ["M1", "M2", "M3"]
    events = [
        ["版本 x.x 发布", "功能 A 上线"],
        ["体验优化", "关键 Bug 修复"],
        ["里程碑交付", "协作项目结项"],
    ]
    line_y = Inches(3.2)
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.2), line_y, Inches(10.9), Pt(3)
    ).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = BLUE_ACCENT
    slide.shapes[-1].line.fill.background()
    for i, month in enumerate(months):
        cx = 2.2 + i * 3.8
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(cx - 0.15), Inches(3.05), Inches(0.3), Inches(0.3)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = BLUE_PRIMARY
        dot.line.fill.background()
        ml = slide.shapes.add_textbox(Inches(cx - 0.4), Inches(2.5), Inches(1), Inches(0.4))
        ml.text_frame.paragraphs[0].text = month
        ml.text_frame.paragraphs[0].font.bold = True
        ml.text_frame.paragraphs[0].font.size = Pt(16)
        ml.text_frame.paragraphs[0].font.color.rgb = BLUE_DARK
        ml.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        bl = slide.shapes.add_textbox(Inches(cx - 1.1), Inches(3.55), Inches(2.2), Inches(2))
        tf = bl.text_frame
        tf.word_wrap = True
        for j, ev in enumerate(events[i]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = f"• {ev}"
            p.font.size = Pt(14)
            p.font.color.rgb = GRAY_DARK
            p.space_after = Pt(8)
    add_footer(slide, prs, footer)


def slide_two_column(prs, title, left_title, left_items, right_title, right_items, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    add_top_bar(slide, prs)
    add_left_accent(slide)
    tb = slide.shapes.add_textbox(Inches(0.95), Inches(0.55), Inches(10), Inches(0.7))
    style_title(tb.text_frame, title, size=32)
    for col, (ctitle, citems, left) in enumerate(
        [(left_title, left_items, 0.6), (right_title, right_items, 6.85)]
    ):
        hdr = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(1.35), Inches(5.9), Inches(0.55)
        )
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = BLUE_PRIMARY if col == 0 else BLUE_DARK
        hdr.line.fill.background()
        ht = hdr.text_frame
        ht.text = ctitle
        ht.paragraphs[0].font.size = Pt(18)
        ht.paragraphs[0].font.color.rgb = WHITE
        ht.paragraphs[0].font.bold = True
        ht.vertical_anchor = MSO_ANCHOR.MIDDLE
        ht.margin_left = Inches(0.2)
        box = slide.shapes.add_textbox(Inches(left + 0.15), Inches(2.05), Inches(5.6), Inches(4.5))
        add_bullets(box.text_frame, citems, size=16)
    add_footer(slide, prs, footer)


def slide_plan(prs, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    add_top_bar(slide, prs)
    add_left_accent(slide)
    tb = slide.shapes.add_textbox(Inches(0.95), Inches(0.55), Inches(8), Inches(0.7))
    style_title(tb.text_frame, "下季度规划", size=32)
    rows = [
        ("目标 O1", "可衡量的目标描述", "负责人"),
        ("目标 O2", "可衡量的目标描述", "负责人"),
        ("目标 O3", "可衡量的目标描述", "负责人"),
    ]
    headers = ["优先级 / 目标", "关键结果 KR", "Owner"]
    col_x = [0.6, 4.2, 9.5]
    col_w = [3.4, 5.1, 2.8]
    hy = 1.4
    for j, h in enumerate(headers):
        cell = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(hy), Inches(col_w[j]), Inches(0.5)
        )
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE_PRIMARY
        cell.line.fill.background()
        cell.text_frame.text = h
        cell.text_frame.paragraphs[0].font.size = Pt(13)
        cell.text_frame.paragraphs[0].font.color.rgb = WHITE
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.margin_left = Inches(0.1)
        cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(rows):
        ry = 1.95 + i * 1.15
        for j, val in enumerate(row):
            cell = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(col_x[j]),
                Inches(ry),
                Inches(col_w[j]),
                Inches(1.0),
            )
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE_LIGHT if i % 2 == 0 else WHITE
            cell.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
            cell.text_frame.text = val
            cell.text_frame.paragraphs[0].font.size = Pt(14)
            cell.text_frame.paragraphs[0].font.color.rgb = GRAY_DARK
            cell.text_frame.margin_left = Inches(0.1)
            cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
    add_footer(slide, prs, footer)


def slide_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_white(slide)
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(2.5), prs.slide_width, Inches(2.5)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE_PRIMARY
    band.line.fill.background()
    thanks = slide.shapes.add_textbox(Inches(0), Inches(3.0), prs.slide_width, Inches(1))
    tp = thanks.text_frame.paragraphs[0]
    tp.text = "感谢聆听"
    tp.font.size = Pt(40)
    tp.font.bold = True
    tp.font.color.rgb = WHITE
    tp.alignment = PP_ALIGN.CENTER
    qa = slide.shapes.add_textbox(Inches(0), Inches(4.0), prs.slide_width, Inches(0.6))
    qp = qa.text_frame.paragraphs[0]
    qp.text = "Q & A"
    qp.font.size = Pt(22)
    qp.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)
    qp.alignment = PP_ALIGN.CENTER
    contact = slide.shapes.add_textbox(Inches(0), Inches(6.2), prs.slide_width, Inches(0.5))
    cp = contact.text_frame.paragraphs[0]
    cp.text = "联系方式：name@company.com"
    cp.font.size = Pt(12)
    cp.font.color.rgb = GRAY_TEXT
    cp.alignment = PP_ALIGN.CENTER


def build_presentation(output: Path) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    footer = "产品季度总结 | 2026 Q__"

    slide_cover(prs)
    slide_toc(prs, footer)
    slide_overview(prs, footer)
    slide_kpi(prs, footer)
    slide_timeline(prs, footer)
    slide_content_list(
        prs,
        "亮点与成果",
        [
            "业务增长：描述核心增长数据与驱动因素",
            "产品创新：本季重要功能或体验升级",
            "用户价值：典型案例或用户反馈摘要",
            "团队协同：跨部门协作成果（可选）",
        ],
        footer,
    )
    slide_two_column(
        prs,
        "问题与挑战",
        "主要问题",
        [
            "问题一：现象与影响",
            "问题二：根因简述",
            "问题三：当前应对状态",
        ],
        "风险与对策",
        [
            "风险：潜在影响范围",
            "对策：已采取 / 计划措施",
            "需支持：资源或决策诉求",
        ],
        footer,
    )
    slide_plan(prs, footer)
    slide_section(prs, "附录", "详细数据、版本清单、用户调研等可放在附录页", footer)
    slide_closing(prs)

    prs.save(output)


if __name__ == "__main__":
    out = Path(__file__).resolve().parent / "产品季度总结模板-简约蓝.pptx"
    build_presentation(out)
    print(f"Saved: {out}")
